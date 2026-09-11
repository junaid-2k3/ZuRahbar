#!/usr/bin/env python3
"""Generate ZuRehbar_Pitch_Deck.pptx using python-pptx.

Usage:
    python scripts/generate_pptx.py
"""

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx is not installed. Installing python-pptx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(7, 9, 14)
    TEXT_WHITE = RGBColor(240, 244, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ACCENT_EMERALD = RGBColor(0, 240, 180)
    ACCENT_CYAN = RGBColor(0, 216, 255)
    ACCENT_CORAL = RGBColor(255, 87, 87)
    CARD_BG = RGBColor(15, 20, 32)
    CARD_BORDER = RGBColor(30, 41, 59)

    slides_data = [
        {
            "tag": "QWEN AI HACKATHON 2026",
            "title": "ZuRehbar (زو رہبر)",
            "subtitle": "Intelligent Voice & Text Navigation for Zu Peshawar Commuters",
            "bullets": [
                "Empowering 250,000+ daily riders on Peshawar BRT (19 total routes).",
                "Deterministic shortest-path transit routing & 9-band distance fare calculation.",
                "Natural multilingual AI support in English, Urdu, Roman Urdu, and Pashto."
            ]
        },
        {
            "tag": "02 • THE PROBLEM",
            "title": "Navigating Zu Peshawar is Hard for Daily Riders",
            "subtitle": "High network complexity coupled with fragmented, non-computable public information.",
            "bullets": [
                "No Trip Planner: TransPeshawar provides static tables but zero point-to-point routing.",
                "Fares Hidden in Urdu JPEGs: Fare tables exist only inside an Urdu image file.",
                "Missing & Conflicting Data: 9 feeder routes lack timetables; card prices contradict on web pages."
            ]
        },
        {
            "tag": "03 • THE SOLUTION",
            "title": "ZuRehbar: Your Conversational Transit Guide",
            "subtitle": "Instant turn-by-turn guidance and fare estimation in any language.",
            "bullets": [
                "Ask Anything: Voice/text queries in English, Urdu, Roman Urdu, or Pashto.",
                "Instant Guidance: Returns exact bus numbers, boarding platforms, transfer stations, and fares.",
                "Offline Resilience: On-device SQLite database (Drift) operates under zero connectivity."
            ]
        },
        {
            "tag": "04 • ARCHITECTURAL CORE",
            "title": "Why Standard LLMs Fail — The Deterministic Split",
            "subtitle": "Preventing LLM transit hallucinations through strict division of responsibility.",
            "bullets": [
                "Principle: The model decides WHICH tool to call; tools decide WHAT IS TRUE.",
                "NetworkX Graph: Computes shortest-path routes, wait times, and transfer nodes.",
                "Qdrant Hybrid Search: Combines Qwen dense vectors + BM25 sparse vectors."
            ]
        },
        {
            "tag": "05 • SYSTEM ARCHITECTURE",
            "title": "End-to-End System Architecture",
            "subtitle": "From raw WP REST scraping & computer vision to Qwen Function Calling and Flutter UI.",
            "bullets": [
                "1. Data Ingestion: Custom HTTP client + AI Vision Tiling for Urdu fare JPEGs.",
                "2. Grounded Engine: Curated JSON -> NetworkX Graph + Qdrant Hybrid Index.",
                "3. Agent Surface: Qwen NLU dispatches plan_journey, get_fare, or search_knowledge."
            ]
        },
        {
            "tag": "06 • GRAPH ENGINE",
            "title": "Solving Transit Math: Route-Stop Graph Modeling",
            "subtitle": "Why standard station graphs fail and how route-stop nodes model wait friction.",
            "bullets": [
                "Route-Stop Nodes: Modeled as (route, direction, station) to capture wait friction.",
                "Weight Formula: Boarding weight = (headway / 2) + 90s transfer penalty.",
                "Distance Interpolation: Leg distance calculated proportionally from travel-time share."
            ]
        },
        {
            "tag": "07 • DATA ENGINEERING",
            "title": "Robust Ingestion from Imperfect Web Sources",
            "subtitle": "Overcoming broken SSL, Urdu image data, and official dataset contradictions.",
            "bullets": [
                "Custom HTTP Client: Bypasses broken TLS certificate chains safely with disk caching.",
                "Urdu Vision Tiling: Slices high-res fare JPEGs into spatial grids for vision extraction.",
                "Honest Anomaly Logging: Data contradictions are recorded in dataset_report.json."
            ]
        },
        {
            "tag": "08 • VECTOR INDEX",
            "title": "Precision Knowledge Retrieval for Non-Trip Queries",
            "subtitle": "Qdrant dense + BM25 sparse hybrid vector search over transit rules and station facts.",
            "bullets": [
                "Qwen Embeddings: Qwen3-Embedding-0.6B (1024-dim) captures multilingual intent.",
                "BM25 Sparse Anchoring: Pinpoints exact station names (e.g. 'Hashtnagri') with 100% precision.",
                "RRF Fusion: Server-side Reciprocal Rank Fusion merges dense and sparse rankings."
            ]
        },
        {
            "tag": "09 • MOBILE APP",
            "title": "Flutter On-Device Mobile Experience",
            "subtitle": "Designed for fast, hands-free commuter interactions on patchy networks.",
            "bullets": [
                "Voice Pipeline: Tap-to-talk mic input with dynamic waveform and barge-in support.",
                "Custom Map Canvas: Schematic metro-style map rendered via Flutter CustomPainter.",
                "Drift SQLite Sync: Bundled local database enables full offline route calculations."
            ]
        },
        {
            "tag": "10 • BUSINESS PITCH",
            "title": "Ready-to-Deploy Companion App for TransPeshawar",
            "subtitle": "High public impact with zero hardware or IT infrastructure overhead.",
            "bullets": [
                "Zero Hardware Cost: Plug-and-play integration with existing static web endpoints.",
                "Reduces Station Congestion: Automated fare and route guidance cuts ticket lines.",
                "Flexible Formats: Official white-label app, WhatsApp chatbot, or station kiosks."
            ]
        },
        {
            "tag": "11 • ROADMAP",
            "title": "Phased Development & Multi-City Expansion",
            "subtitle": "From Qwen Hackathon MVP to national urban transit AI assistant.",
            "bullets": [
                "Phase 1 & 2 (MVP): Scraper pipeline, NetworkX graph, Qdrant index, 69 pytest suite, Flutter client.",
                "Phase 3 (GTFS/GPS): Map stop GPS coordinates for real-world spatial matching.",
                "Phase 4 (Expansion): Multi-city expansion to Lahore, Islamabad, and Karachi BRT networks."
            ]
        },
        {
            "tag": "12 • CONCLUSION",
            "title": "Transforming Public Mobility with AI",
            "subtitle": "ZuRehbar — Guidance for Every Commuter.",
            "bullets": [
                "Team: Muhammad Hamza, Junaid Ahmad, Sajid Islam & Team.",
                "Live Demo Script: python scripts/ask.py --demo",
                "Contact & Repository: team@zurehbar.pk"
            ]
        }
    ]

    for slide_info in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Background Fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK

        # Header Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = slide_info["tag"]
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_EMERALD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_info["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.5))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = slide_info["subtitle"]
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = ACCENT_CYAN

        # Content Card Box
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.6), Inches(11.7), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Bullets
        content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.9), Inches(10.9), Inches(3.6))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for i, bullet in enumerate(slide_info["bullets"]):
            p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.name = "Arial"
            p.font.size = Pt(18)
            p.font.color.rgb = TEXT_WHITE
            p.space_after = Pt(20)

    output_path = Path(__file__).resolve().parent.parent / "ZuRehbar_Pitch_Deck.pptx"
    prs.save(str(output_path))
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
